#!/usr/bin/env python3
"""CN-Codex 技术讲解 PPT 生成脚本"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ============================================================
# 设计系统常量（匹配 CN-Codex DESIGN.md）
# ============================================================
ACCENT = RGBColor(0x22, 0xC5, 0x5E)       # 翡翠绿
ACCENT_STRONG = RGBColor(0x4A, 0xDE, 0x80)
ACCENT_SOFT = RGBColor(0x1A, 0x3A, 0x2A)
BG_DARK = RGBColor(0x1A, 0x1A, 0x1A)       # 深色背景
BG_SURFACE = RGBColor(0x26, 0x26, 0x26)    # 二级表面
BG_CARD = RGBColor(0x2A, 0x2A, 0x2A)       # 卡片
TEXT_STRONG = RGBColor(0xF0, 0xF0, 0xF0)   # 标题/强调
TEXT_BASE = RGBColor(0xC8, 0xC8, 0xC8)     # 正文
TEXT_MUTED = RGBColor(0x88, 0x88, 0x88)    # 次要
TEXT_FAINT = RGBColor(0x66, 0x66, 0x66)    # 微弱
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

# ============================================================
# 工具函数
# ============================================================
def add_bg(slide, color=BG_DARK):
    """设置幻灯片背景"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color=None, line_color=None, shape_type=MSO_SHAPE.RECTANGLE):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_textbox(slide, left, top, width, height, text="", font_size=14, color=TEXT_BASE, bold=False, alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_paragraph(text_frame, text, font_size=14, color=TEXT_BASE, bold=False, space_before=Pt(4), space_after=Pt(2), alignment=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.space_before = space_before
    p.space_after = space_after
    p.alignment = alignment
    return p

def add_accent_bar(slide, left=Inches(0), top=Inches(0.5), width=Inches(0.08), height=Inches(0.4)):
    bar = add_shape(slide, left, top, width, height, fill_color=ACCENT)
    return bar

def make_section_slide(title, subtitle=""):
    """章节标题页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_bg(slide, BG_DARK)
    # 顶部装饰条
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), fill_color=ACCENT)
    # 左侧竖线
    add_shape(slide, Inches(1.5), Inches(2.5), Inches(0.06), Inches(2.5), fill_color=ACCENT)
    # 标题
    add_textbox(slide, Inches(2.0), Inches(2.8), Inches(9), Inches(1.2), 
                title, font_size=36, color=TEXT_STRONG, bold=True)
    if subtitle:
        add_textbox(slide, Inches(2.0), Inches(4.0), Inches(9), Inches(0.8), 
                    subtitle, font_size=18, color=TEXT_MUTED)
    return slide

def make_content_slide(title, bullets, accent_color=ACCENT):
    """内容页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_DARK)
    # 顶部装饰条
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.04), fill_color=accent_color)
    # 标题
    add_accent_bar(slide, Inches(0.4), Inches(0.4))
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11), Inches(0.6),
                title, font_size=26, color=TEXT_STRONG, bold=True)
    # 分割线
    add_shape(slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.015), fill_color=RGBColor(0x33, 0x33, 0x33))
    # 内容
    y_pos = Inches(1.2)
    for bullet in bullets:
        if isinstance(bullet, tuple):
            text, level = bullet
            left = Inches(0.8 + level * 0.4)
            fs = 16 if level == 0 else 14
            clr = TEXT_STRONG if level == 0 else TEXT_BASE
        else:
            text = bullet
            left = Inches(0.8)
            fs = 16
            clr = TEXT_STRONG
        
        tb = add_textbox(slide, left, y_pos, Inches(11.5 - (left - Inches(0.8))), Inches(0.4),
                         text, font_size=fs, color=clr)
        y_pos += Inches(0.42)
    return slide

def make_code_slide(title, code_text):
    """代码展示页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_DARK)
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.04), fill_color=ACCENT)
    add_accent_bar(slide, Inches(0.4), Inches(0.4))
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11), Inches(0.6),
                title, font_size=24, color=TEXT_STRONG, bold=True)
    # 代码块背景
    box = add_shape(slide, Inches(0.6), Inches(1.1), Inches(12), Inches(5.8), 
                    fill_color=RGBColor(0x1E, 0x1E, 0x1E), 
                    line_color=RGBColor(0x33, 0x33, 0x33))
    # 代码文本
    tb = add_textbox(slide, Inches(0.9), Inches(1.3), Inches(11.5), Inches(5.5),
                     code_text, font_size=12, color=TEXT_BASE, font_name="Consolas")
    return slide

def make_two_col_slide(title, left_text, right_text):
    """双栏内容页"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, BG_DARK)
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.04), fill_color=ACCENT)
    add_accent_bar(slide, Inches(0.4), Inches(0.4))
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11), Inches(0.6),
                title, font_size=24, color=TEXT_STRONG, bold=True)
    # 左栏
    add_textbox(slide, Inches(0.6), Inches(1.2), Inches(5.8), Inches(5.5),
                left_text, font_size=14, color=TEXT_BASE)
    # 右栏
    add_textbox(slide, Inches(6.8), Inches(1.2), Inches(5.8), Inches(5.5),
                right_text, font_size=14, color=TEXT_BASE)
    return slide

# ============================================================
# 创建幻灯片
# ============================================================

# ---------- 封面 ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG_DARK)
# 背景装饰
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), fill_color=ACCENT)
add_shape(slide, Inches(0), Inches(7.42), SLIDE_W, Inches(0.08), fill_color=ACCENT)
# 中心内容
add_textbox(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(1.2),
            "CN-Codex", font_size=54, color=ACCENT_STRONG, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(2.8), Inches(10), Inches(0.8),
            "AI 驱动的全栈编程助手桌面应用", font_size=28, color=TEXT_STRONG, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(3.6), Inches(10), Inches(0.6),
            "技术架构详解 · 核心模块 · 实现原理", font_size=18, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
# 技术标签
tags = ["Tauri v2", "React 18", "Rust", "TypeScript", "DeepSeek"]
x_start = Inches(3.5)
for i, tag in enumerate(tags):
    tag_box = add_shape(slide, x_start + Inches(i * 1.6), Inches(4.6), Inches(1.4), Inches(0.45),
                         fill_color=BG_CARD, line_color=ACCENT)
    tag_box.line.width = Pt(1)
    add_textbox(slide, x_start + Inches(i * 1.6), Inches(4.63), Inches(1.4), Inches(0.4),
                tag, font_size=12, color=ACCENT, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(5.6), Inches(10), Inches(0.5),
            "版本 0.1.0  |  基于 Tauri 2 构建  |  已在 DeepSeek v4 Flash 模型下验证",
            font_size=13, color=TEXT_FAINT, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(6.3), Inches(10), Inches(0.5),
            "团队内部分享 · 2025",
            font_size=14, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)

# ---------- 目录 ----------
make_content_slide("目录", [
    "1. 项目背景与定位",
    "2. 技术栈全景",
    "3. 整体架构设计",
    "4. 前端架构详解",
    "5. Rust 后端架构详解",
    "6. 核心引擎 AgentEngine",
    "7. 工具执行器 ToolExecutor",
    "8. LLM 适配层",
    "9. 数据持久化",
    ("10. 移动端远程控制", 0),
    ("11. 外部浏览器与录制回放", 0),
    ("12. SmartBrain 智能脑", 0),
    ("13. 机器人系统", 0),
    ("14. 构建与发布", 0),
])

# ========== 第1章：项目背景与定位 ==========
make_section_slide("01", "项目背景与定位")

make_content_slide("核心定位", [
    ("桌面原生应用", 0),
    "  基于 Tauri 2 框架，提供原生桌面体验",
    ("AI 编程助手", 0),
    "  支持多轮对话、工具调用、Goal 模式，自动完成编程任务",
    ("手机远程控制", 0),
    "  扫码连接手机，随时随地监控和控制 AI 助手",
    ("机器人自动化", 0),
    "  AI 自动创建专业角色，绑定技能和工作流节点，实现自动化流水线",
    ("全流程开发支持", 0),
    "  从项目创建、代码编写、测试到部署的全流程覆盖",
])

make_content_slide("核心亮点", [
    ("🔹 手机远程控制", 0),
    "  实时监控 · 远程操作 · 双重连接（局域网直连 + 公网中转）",
    ("🔹 机器人系统", 0),
    "  智能角色创建 · 技能绑定 · 工作流编排 · 多机器人协作",
    ("🔹 全栈开发流程", 0),
    "  项目创建 → 代码开发 → 自动测试 → 持续迭代",
    ("🔹 SmartBrain 智能脑", 0),
    "  经验自动提取 · 知识库管理 · BM25 搜索 · 自动注入系统提示词",
    ("🔹 操作录制与回放", 0),
    "  外部浏览器 CDP 协议 · 实时录制 · 智能回放",
])

# ========== 第2章：技术栈全景 ==========
make_section_slide("02", "技术栈全景")

make_two_col_slide("前端技术栈", 
    "▎桌面框架：Tauri v2\n" +
    "▎UI 框架：React 18 + TypeScript\n" +
    "▎构建工具：Vite 6\n" +
    "▎状态管理：Zustand 5\n" +
    "▎样式框架：Tailwind CSS 4\n" +
    "▎国际化：react-intl\n" +
    "▎Markdown：react-markdown 9\n" +
    "▎代码高亮：rehype-highlight\n" +
    "▎图标库：@tabler/icons-react\n" +
    "▎终端：@xterm/xterm 6",
    "▎测试框架：Vitest + jsdom\n" +
    "▎多页面入口：3个独立窗口\n" +
    "  - main（主聊天窗口）\n" +
    "  - detail（文档详情窗口）\n" +
    "  - diff（差异对比窗口）"
)

make_two_col_slide("后端技术栈 (Rust)", 
    "▎系统语言：Rust 1.85+ (Edition 2024)\n" +
    "▎桌面框架：Tauri 2\n" +
    "▎异步运行时：tokio (full)\n" +
    "▎HTTP 客户端：reqwest\n" +
    "▎序列化：serde / serde_json\n" +
    "▎数据库：rusqlite (SQLite bundled)\n" +
    "▎配置解析：toml 0.8\n" +
    "▎移动端服务器：axum 0.8\n" +
    "▎WebSocket：tokio-tungstenite",
    "▎日志：tracing\n" +
    "▎OCR：ndarray + ort (ONNX Runtime)\n" +
    "▎终端：portable-pty\n" +
    "▎文档解析：pdf-extract, calamine\n" +
    "▎QR 码：qrcode\n" +
    "▎图片处理：image\n" +
    "▎Tauri 插件：\n" +
    "  - shell / dialog / opener\n" +
    "  - fs / process / notification"
)

# ========== 第3章：整体架构 ==========
make_section_slide("03", "整体架构设计")

make_content_slide("三层次架构", [
    ("前端层 (React + TypeScript)", 0),
    "  UI 渲染、状态管理、事件监听、用户交互",
    ("中间层 (Tauri IPC)", 0),
    "  invoke/events 双向通信、JSON-RPC 协议",
    ("后端层 (Rust)", 0),
    "  Agent 引擎、工具执行、LLM 适配、数据持久化",
])

make_content_slide("通信模型与事件流", [
    ("通信模型", 0),
    "  前端 (React) ←→ Tauri IPC ←→ Rust 后端 ←→ LLM API (HTTP SSE)",
    ("事件流（流式渲染）", 0),
    "  LLM API → SSE Stream → Rust Agent → Tauri emit()",
    "  → 前端 listener → Zustand Store → React 重渲染",
    ("关键设计原则", 0),
    "  • 深色优先：IDE 原生暗色界面，翡翠绿强调色",
    "  • 紧凑高效：中高密度布局，去阴影化设计",
    "  • 中英文双语：默认中文，支持运行时切换",
])

# ========== 第4章：前端架构 ==========
make_section_slide("04", "前端架构详解")

make_content_slide("React 组件树", [
    ("App.tsx（根组件）", 0),
    "  ├─ ErrorBoundary / IntlProvider",
    "  ├─ TitleBar（自定义无边框标题栏）",
    "  ├─ Sidebar（左侧：会话列表 + 项目分组）",
    "  ├─ ChatPage（主聊天区）",
    "  │   ├─ MessageList（流式消息列表）",
    "  │   ├─ ChatInput（智能输入框 + 斜杠命令）",
    "  │   └─ StreamingIndicator",
    "  ├─ RightPanel（右侧：终端/浏览器/Git）",
    "  └─ StatusBar（底部状态栏）",
    "  模态层：SettingsPanel / ApprovalModal / 录制定位浮窗",
])

make_content_slide("Zustand 状态管理", [
    ("appStore.ts（~1300行）— 主状态", 0),
    "  • 会话管理：threads, messages, streamingText",
    "  • 项目管理：projects, currentProjectId, workspaceCwd",
    "  • 模型配置：providers, activeModelId, configuredModels",
    "  • 布局状态：sidebarWidth, rightPanelWidth, resizer",
    ("settingsStore.ts — 设置状态", 0),
    "  • 语言切换：zh-CN / en-US",
    "  • 主题切换：dark / light / system",
    ("useTauriEvents.ts — 事件桥接", 0),
    "  后端 Tauri emit → Zustand store → React 重渲染",
])

make_content_slide("API 层封装（invoke 调用）", [
    ("13 个 API 模块，封装所有 Tauri Command", 0),
    "  standalone.ts — 核心引擎（thread/chat/config）",
    "  git.ts — Git 操作（status/diff/commit/push 等 13 个命令）",
    "  window.ts — 窗口控制（主窗/浏览器/详情/Diff 窗口）",
    "  app_state.ts — SQLite KV 持久化",
    "  approval.ts / fileReview.ts — 审批与文件审阅",
    "  usage.ts — 用量统计查询",
    "  hook.ts / plugin.ts / robot.ts / skill.ts",
    "  recording.ts / workflow.ts",
])

# ========== 第5章：Rust 后端 ==========
make_section_slide("05", "Rust 后端架构")

make_content_slide("后端模块结构（27 个核心模块）", [
    ("模块分类", 0),
    "  核心引擎：agent.rs, subagent_engine.rs, tool_executor.rs",
    "  LLM 适配：adapter/（4 种 API 格式）",
    "  数据持久化：thread_store.rs, usage/db.rs",
    "  配置管理：config_system.rs, state.rs",
    "  移动端：mobile_server.rs（Axum + WebSocket）",
    "  浏览器：external_browser.rs, recording.rs",
    "  AI 增强：smartbrain/, compaction.rs, hook_runtime.rs",
    "  机器人：robot_orchestrator.rs, robot_loader.rs",
    "  插件：plugin_loader.rs",
    "  工具：ocr/, terminal.rs, document_parser.rs, git_service.rs",
])

make_content_slide("AppState — 应用全局状态", [
    ("Rust 端全局状态管理", 0),
    "  locale — 语言（默认 zh-CN）",
    "  current_thread_id — 当前活跃会话",
    "  project_root — 项目根目录",
    "  workspace_config_dir — codey/ 配置目录",
    "  config_manager — ConfigManager（TOML 读写）",
    "  thread_store — Arc<ThreadStore>（会话 JSONL 持久化）",
    "  agent_engine — Arc<AgentEngine>（Agent 引擎）",
    "  usage_db — Arc<UsageDb>（用量 SQLite 数据库）",
    "  approval_tx/rx — 审批通道（mpsc）",
    "  external_browser — 外部浏览器管理",
    "  recorder — 录制引擎",
])

# ========== 第6章：AgentEngine ==========
make_section_slide("06", "核心引擎 AgentEngine")

make_content_slide("AgentEngine 运行流程", [
    ("run_turn(thread_id, user_input, mode) 执行流程", 0),
    "  1. 解析模式和配置",
    "  2. 构建消息历史（InternalMessage 格式）",
    "  3. 调用 LLM API（流式 SSE）",
    "  4. 解析响应（文本 / ToolCalls）",
    "  5. 执行工具调用（ToolExecutor）",
    "  6. 收集工具结果，注入上下文",
    "  7. 重复 3-6 直到完成",
    ("支持模式", 0),
    "  chat — 普通聊天模式",
    "  goal — 目标驱动模式（自动循环）",
    "  plan — 计划模式",
    "  robot-create / robot-modify — 机器人创建/修改模式",
])

make_content_slide("Goal 模式自动循环", [
    ("用户设置目标 → AI 自主执行", 0),
    "  Goal → LLM 调用 → 工具执行 → 检查完成条件",
    "  → 未完成则继续 → 达到目标／Token 预算耗尽 → 完成",
    ("关键特性", 0),
    "  • 自动中断：用户可随时点击停止按钮",
    "  • Token 预算：支持 goal_budget_tokens 限制",
    "  • 上下文压缩：超阈值时自动触发 /compact",
    "  • 文件变更跟踪：记录每个 turn 的增删改文件",
    "  • RunSummary：每次 turn 结束后生成变更摘要",
])

# ========== 第7章：ToolExecutor ==========
make_section_slide("07", "工具执行器 ToolExecutor")

make_content_slide("40+ 内置工具分类", [
    ("核心开发工具", 0),
    "  shell/shell_command, exec_command, read/write_file",
    "  apply_patch, list_directory, code_review",
    ("浏览器工具", 0),
    "  browser_run, view_image, ocr_image, image_generate",
    ("AI 辅助工具", 0),
    "  update_plan, request_user_input, tool_search",
    ("网络工具", 0),
    "  web_search（双引擎：DuckDuckGo + Bing）, web_fetch",
    ("MCP 工具", 0),
    "  mcp_call_tool（动态调用）+ mcp__*（直接调用）",
    ("记忆/知识工具", 0),
    "  memory_read/write/search/update/forget/list",
    ("子代理工具", 0),
    "  spawn_agent, wait_agent, send_input, close_agent",
])

# ========== 第8章：LLM 适配层 ==========
make_section_slide("08", "LLM 适配层")

make_content_slide("4 种 API 格式适配", [
    ("ProviderAdapter Trait", 0),
    "  build_url() — 构建请求 URL",
    "  build_headers() — 构建请求头",
    "  build_body() — 构建请求体",
    "  parse_stream_line() — 解析 SSE 行",
    ("支持的 API 格式", 0),
    "  chat (OpenAI Chat Completions) — 默认，最通用",
    "  responses (OpenAI Responses API) — OpenAI 原生",
    "  anthropic (Anthropic Messages API) — Anthropic Claude",
    "  gemini (Google Gemini API) — Google Gemini",
])

make_content_slide("流式处理架构", [
    ("统一流式事件 StreamEvent", 0),
    "  TextDelta(String) — 文本增量",
    "  ToolCallDelta { index, id, name, arguments }",
    "  Done { finish_reason } — 流结束",
    "  Usage(UsageInfo) — Token 用量",
    ("适配器工作流程", 0),
    "  HTTP POST → SSE Stream → parse_stream_line()",
    "  → StreamEvent → AgentEngine 累积拼接",
    "  → 完整文本 / 完整 ToolCall → 分发执行",
])

make_code_slide("适配器代码结构", 
    "// adpater/mod.rs — 工厂方法\n" +
    "pub fn get_adapter(wire_api: &str) -> Box<dyn ProviderAdapter> {\n" +
    "    match wire_api {\n" +
    '        "responses" => Box::new(responses::ResponsesAdapter),\n' +
    '        "anthropic" => Box::new(anthropic::AnthropicAdapter),\n' +
    '        "gemini" => Box::new(google::GoogleAdapter),\n' +
    '        _ => Box::new(chat_completions::ChatCompletionsAdapter),\n' +
    "    }\n" +
    "}\n\n" +
    "// types.rs — 统一类型定义\n" +
    "pub struct InternalMessage {\n" +
    "    pub role: String,\n" +
    "    pub content: Option<Value>,\n" +
    "    pub tool_calls: Option<Vec<InternalToolCall>>,\n" +
    "    pub tool_call_id: Option<String>,\n" +
    "}")

# ========== 第9章：数据持久化 ==========
make_section_slide("09", "数据持久化")

make_content_slide("三层持久化体系", [
    ("第1层：ThreadStore — 会话存储（JSONL）", 0),
    "  每次 turn 写入 JSONL 文件",
    "  支持 goal 状态、机器人状态持久化",
    "  支持上下文压缩（compaction）",
    ("第2层：SQLite — 用量统计 + KV 存储", 0),
    "  usage.db：记录每次 LLM 调用的 Token 用量和费用",
    "  app_state：前端持久化状态（项目/模型/布局）",
    ("第3层：TOML — 用户配置", 0),
    "  config.toml：模型配置、供应商、MCP 服务器、Hook",
])

# ========== 第10章：移动端远程控制 ==========
make_section_slide("10", "移动端远程控制")

make_content_slide("Axum Web 服务器 + WebSocket 架构", [
    ("手机扫码连接 AI 助手", 0),
    ("服务器组件", 0),
    "  • Axum HTTP 服务器（0.0.0.0:动态端口）",
    "  • WebSocket 实时推送（Tauri emit → 广播到所有移动端）",
    "  • QR 码生成（qrcode crate，200×200 SVGs）",
    ("API 端点", 0),
    "  GET  /api/threads — 列出所有会话",
    "  GET  /api/threads/{id}/messages — 获取消息",
    "  POST /api/threads/{id}/chat — 发送消息",
    "  POST /api/threads/{id}/interrupt — 中断",
    "  WS  /ws — WebSocket 实时推送（事件广播）",
    ("连接方式", 0),
    "  • 局域网直连（低延迟）",
    "  • 公网中转服务器（跨网络访问）",
])

# ========== 第11章：外部浏览器与录制 ==========
make_section_slide("11", "外部浏览器与录制回放")

make_content_slide("Chrome CDP 协议集成", [
    ("ExternalBrowser 管理器", 0),
    "  自动检测 Chrome/Edge 安装路径",
    "  启动 `--remote-debugging-port=9222`",
    "  使用临时 user-data-dir 隔离",
    "  CDP 就绪检测（15 秒超时）",
    ("浏览器发现路径（Windows）", 0),
    "  C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe",
    "  C:\\Program Files (x86)\\Google\\Chrome\\Application\\chrome.exe",
    "  %LOCALAPPDATA%\\Google\\Chrome\\Application\\chrome.exe",
    "  Microsoft Edge 作为备选",
])

make_content_slide("操作录制引擎", [
    ("CDP 注入 JS 脚本录制用户操作", 0),
    "  Runtime.addBinding('__rr_push')",
    "  JS 注入 → 捕获 click/input/submit/change/navigate",
    "  通过 WebSocket 实时推送事件到 Rust",
    ("录制流程", 0),
    "  start_recording() → 注入 JS → CDP reader 后台监听",
    "  → 事件累积 → stop_recording() → 生成 TraceFile",
    ("TraceFile 格式", 0),
    "  { session_id, events: [{ type, timestamp, url, selector }] }",
    ("智能回放", 0),
    "  支持录制的操作序列自动回放执行",
])

# ========== 第12章：SmartBrain ==========
make_section_slide("12", "SmartBrain 智能脑")

make_content_slide("启动流水线", [
    ("应用启动时自动执行 4 阶段任务", 0),
    ("阶段1：经验提取（Extractor）", 1),
    "  从历史会话中提取经验 → memories/experiences/raw/",
    "  每次启动最多处理 5 条未提取的会话",
    ("阶段2：经验合并（Consolidator）", 1),
    "  合并相似经验 → 精简存储",
    ("阶段3：知识扫描（Knowledge Scanner）", 1),
    "  扫描并索引知识文档 → 构建 BM25 倒排索引",
    ("阶段4：索引重建（Index Builder）", 1),
    "  生成根索引文档 → 准备全文搜索",
])

make_content_slide("BM25 搜索与知识管理", [
    ("搜索架构", 0),
    "  smartbrain_search(query) → BM25 全文检索",
    "  支持按概念类型、标签、来源类型筛选",
    "  结果自动注入到 Agent 的 system prompt",
    ("知识管理命令", 0),
    "  smartbrain_list_knowledge — 列出知识文档",
    "  smartbrain_read_knowledge — 读取知识内容",
    "  smartbrain_upload_knowledge — 上传知识文档",
    "  smartbrain_search — BM25 全文搜索",
    "  smartbrain_rebuild_index — 重建搜索索引",
    ("OKF 格式（Open Knowledge Format）", 0),
    "  统一 Frontmatter 格式存储经验和知识",
])

# ========== 第13章：机器人系统 ==========
make_section_slide("13", "机器人系统")

make_content_slide("机器人工作流编排", [
    ("RobotOrchestrator 核心流程", 0),
    "  prepare_state() → 编译 workflowNodes 为 runtime_nodes",
    "  → bind_goal_to_current_node() → 注入 overlay prompt",
    "  → Agent 执行当前节点 → 检测 <workflow_node_done/> 标记",
    "  → advance_node() → 推进到下一节点",
    ("关键技术细节", 0),
    "  • 节点推进标记：<workflow_node_done/>（控制 token）",
    "  • Overlay Prompt：补充 goal 模式，不覆盖 system prompt",
    "  • 模板变量：{{goal}} / {{objective}} 动态替换",
    ("内置机器人示例", 0),
    "  requirements-design-bot — 需求分析机器人",
    "  fullstack-dev-bot — 全栈开发机器人",
    "  qa-test-bot — 测试机器人",
])

# ========== 第14章：构建与发布 ==========
make_section_slide("14", "构建与发布流程")

make_content_slide("三种构建模式", [
    ("开发模式 (dev.bat)", 0),
    "  检查 Vite 端口 1420 → 确保 Node.js 便携版 → Tauri dev",
    "  （Vite HMR + Rust 热重载）",
    ("构建模式 (build.bat)", 0),
    "  pnpm build → pnpm tauri build → 复制到 build/",
    ("发布模式 (publish.bat)", 0),
    "  1. 清理 publish/ 目录",
    "  2. 安装依赖 + 构建 mobile-web",
    "  3. Tauri Release 构建（--no-bundle 便携版）",
    "  4. 复制 CN-Codex.exe + DLL + 运行时资源",
    "  5. 打包 Node.js 便携版（v22.16.0）",
    "  6. 生成可分发的 publish/ 目录",
])

# ========== 第15章：总结 ==========
make_section_slide("15", "总结与展望")

make_content_slide("项目亮点总结", [
    ("技术特色", 0),
    "  • 全栈 Rust + React 架构，性能优异的桌面原生应用",
    "  • 4 种 LLM API 格式适配，灵活切换供应商",
    "  • 40+ 内置工具，覆盖文件/命令/浏览器/MCP/子代理",
    "  • 内置离线 OCR（PP-OCRv5 Mobile），无需网络",
    ("创新功能", 0),
    "  • 手机远程控制（扫码即连，实时监控）",
    "  • 机器人系统（AI 自动创建专业角色 + 工作流编排）",
    "  • SmartBrain 智能脑（经验自动提取 + BM25 知识检索）",
    "  • 外部浏览器录制回放（CDP 协议）",
])

make_content_slide("技术数据", [
    ("项目规模", 0),
    "  • 前端：~15,000 行 TypeScript/React 代码",
    "  • 后端：~25,000 行 Rust 代码",
    "  • 模块数：27 个核心 Rust 模块",
    "  • API 命令：80+ Tauri Commands",
    "  • 内置工具：40+",
    "  • 内置插件：8 个",
    "  • 内置 Skills：60+",
    ("测试覆盖", 0),
    "  • Rust：内联测试 + 集成测试，覆盖所有核心模块",
    "  • 前端：Vitest + jsdom 单元测试",
])

# ---------- 尾页 ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BG_DARK)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), fill_color=ACCENT)
add_shape(slide, Inches(0), Inches(7.42), SLIDE_W, Inches(0.08), fill_color=ACCENT)
add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.0),
            "Thank You", font_size=48, color=ACCENT_STRONG, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(3.2), Inches(10), Inches(0.8),
            "手机在手，代码我有", font_size=24, color=TEXT_STRONG, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(4.2), Inches(10), Inches(0.6),
            "https://github.com/longdream/cn-codex", font_size=16, color=TEXT_MUTED, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(5.5), Inches(10), Inches(0.5),
            "Q & A", font_size=28, color=TEXT_BASE, alignment=PP_ALIGN.CENTER)

# ============================================================
# 保存
# ============================================================
output_path = "D:\\rustwork\\cn-codex\\docs\\CN-Codex-技术讲解.pptx"
prs.save(output_path)
print(f"PPTX 已生成: {output_path}")
print(f"共 {len(prs.slides)} 页幻灯片")